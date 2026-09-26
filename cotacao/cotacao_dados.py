"""Dados da cotação: leitura da planilha COTAÇÃO COMPLETA, das respostas
dos fornecedores e cálculo do comparativo (mesma lógica da aba COMPRAS).

Não depende de interface gráfica, então pode ser usado direto pelo sistema.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Diferença entre o menor e o 2º menor preço abaixo da qual a linha fica
# destacada (na planilha: formatação condicional da coluna PORCENTAGEM < 2%).
LIMITE_DIFERENCA = 0.02


@dataclass
class Item:
    seq: str
    codigo: str
    similar: str
    marca_exigida: str
    qtd: float
    descricao: str

    @property
    def nr_fabrica(self) -> str:
        # Mesmo texto que a macro SIMILAR monta: CÓDIGO {SIMILAR}
        return f"{self.codigo} {{{self.similar}}}" if self.similar else self.codigo


@dataclass
class Preco:
    valor: float
    marca: str = ""


@dataclass
class Cotacao:
    itens: list[Item] = field(default_factory=list)
    fornecedores: list[str] = field(default_factory=list)
    # precos[fornecedor][seq] = Preco
    precos: dict[str, dict[str, Preco]] = field(default_factory=dict)
    # Fornecedores que já mandaram resposta (coluna CHECK da aba FORNECEDORES)
    respondidos: set[str] = field(default_factory=set)

    def adicionar_fornecedor(self, nome: str) -> str:
        nome = nome.strip().upper()
        if nome and nome not in self.fornecedores:
            self.fornecedores.append(nome)
        return nome

    def definir_preco(self, fornecedor: str, seq: str, valor, marca: str = "") -> None:
        v = converter_valor(valor)
        tabela = self.precos.setdefault(fornecedor, {})
        if v is None:
            tabela.pop(seq, None)
        else:
            tabela[seq] = Preco(v, (marca or "").strip())

    def preco(self, fornecedor: str, seq: str) -> Preco | None:
        return self.precos.get(fornecedor, {}).get(seq)

    def comparativo(self, item: Item) -> dict:
        """MENOR / EMPRESA / 2º MENOR / EMPRESA / PORCENTAGEM de um item."""
        ofertas = sorted(
            (p.valor, f) for f in self.fornecedores
            if (p := self.preco(f, item.seq)) is not None
        )
        r = {"menor": None, "empresa": "", "segundo": None, "empresa2": "", "diferenca": None}
        if ofertas:
            r["menor"], r["empresa"] = ofertas[0]
        if len(ofertas) > 1:
            r["segundo"], r["empresa2"] = ofertas[1]
            if r["menor"]:
                r["diferenca"] = r["segundo"] / r["menor"] - 1
        return r

    def total_melhor_preco(self) -> float:
        total = 0.0
        for it in self.itens:
            m = self.comparativo(it)["menor"]
            if m is not None:
                total += m * (it.qtd or 0)
        return total


# ---------------------------------------------------------------- leitura

def converter_valor(valor) -> float | None:
    """Aceita 12.5, '12,50', 'R$ 1.234,56'. Vazio, zero ou texto = sem preço."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor) if valor > 0 else None
    s = re.sub(r"[^\d,.\-]", "", str(valor))
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        v = float(s)
    except ValueError:
        return None
    return v if v > 0 else None


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _normalizar_codigo(c) -> str:
    return re.sub(r"\s+", "", _texto(c)).upper()


def _numero(v, padrao=1.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return padrao


def carregar_planilha(caminho: str) -> Cotacao:
    """Lê o arquivo COTAÇÃO COMPLETA.xlsm.

    Itens vêm da aba PLANILHA (a cotação enviada). Fornecedores vêm do
    cabeçalho da aba COMPRAS (cada fornecedor ocupa duas colunas: valor e
    marca). Preços já digitados na aba COMPRAS também são carregados.
    """
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    cot = Cotacao()

    if "PLANILHA" in wb.sheetnames:
        for linha in wb["PLANILHA"].iter_rows(min_row=6, max_col=8, values_only=True):
            seq, codigo, similar, marca, qtd, desc = (list(linha) + [None] * 8)[:6]
            if not _texto(codigo):
                continue
            cot.itens.append(Item(
                seq=_texto(seq) or str(len(cot.itens) + 1),
                codigo=_texto(codigo),
                similar=_texto(similar),
                marca_exigida=_texto(marca),
                qtd=_numero(qtd),
                descricao=_texto(desc),
            ))

    colunas_fornecedor: dict[str, int] = {}
    if "COMPRAS" in wb.sheetnames:
        ws = wb["COMPRAS"]
        cabecalho = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        # Colunas F em diante, até a primeira coluna de resultado (vazia/AB).
        for i in range(5, len(cabecalho)):
            nome = _texto(cabecalho[i]).upper()
            if not nome or nome == "PORCENTAGEM":
                if nome == "PORCENTAGEM" or colunas_fornecedor:
                    break
                continue
            if nome not in colunas_fornecedor:
                colunas_fornecedor[nome] = i
                cot.fornecedores.append(nome)

        # Itens já montados na aba COMPRAS (caso a PLANILHA esteja vazia)
        # e preços já digitados.
        por_seq = {it.seq: it for it in cot.itens}
        for linha in ws.iter_rows(min_row=2, values_only=True):
            linha = list(linha)
            seq = _texto(linha[0] if linha else None)
            nr = _texto(linha[1] if len(linha) > 1 else None)
            if not seq and not nr:
                continue
            seq = seq or str(len(por_seq) + 1)
            if seq not in por_seq:
                m = re.match(r"^(.*?)\s*\{(.*)\}\s*$", nr)
                codigo, similar = (m.group(1), m.group(2)) if m else (nr, "")
                it = Item(seq, codigo.strip(), similar.strip(),
                          _texto(linha[2] if len(linha) > 2 else None),
                          _numero(linha[3] if len(linha) > 3 else None),
                          _texto(linha[4] if len(linha) > 4 else None))
                cot.itens.append(it)
                por_seq[seq] = it
            for nome, col in colunas_fornecedor.items():
                valor = linha[col] if col < len(linha) else None
                marca = linha[col + 1] if col + 1 < len(linha) else None
                if converter_valor(valor) is not None:
                    cot.definir_preco(nome, seq, valor, _texto(marca))
                    cot.respondidos.add(nome)

    wb.close()
    return cot


def ler_resposta(caminho: str) -> list[dict]:
    """Lê o arquivo que o fornecedor devolveu (o mesmo gerado pelo EXPORTAR:
    SEQ, CÓDIGO DO PRODUTO, ..., VALOR, MARCA). Procura o cabeçalho sozinho,
    então funciona mesmo se o fornecedor mexeu nas linhas de cima.
    Aceita .xlsx, .xlsm e .csv.
    """
    if caminho.lower().endswith(".csv"):
        import csv
        with open(caminho, encoding="utf-8-sig", newline="") as f:
            amostra = f.read(4096)
            f.seek(0)
            dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t")
            linhas = list(csv.reader(f, dialeto))
    else:
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()

    idx = {}
    inicio = None
    for n, linha in enumerate(linhas[:20]):
        nomes = [_texto(c).upper() for c in linha]
        if "VALOR" in nomes and any(x.startswith("CÓDIGO") or x.startswith("CODIGO") for x in nomes):
            for i, nome in enumerate(nomes):
                if nome.startswith("CÓDIGO") or nome.startswith("CODIGO"):
                    idx.setdefault("codigo", i)
                elif nome == "SEQ":
                    idx["seq"] = i
                elif nome == "VALOR":
                    idx["valor"] = i
                elif nome == "MARCA":
                    idx["marca"] = i
            inicio = n + 1
            break
    if inicio is None:
        raise ValueError(
            "Não achei as colunas CÓDIGO e VALOR nesse arquivo. "
            "Use o arquivo de cotação que foi enviado ao fornecedor.")

    def col(linha, chave):
        i = idx.get(chave)
        return linha[i] if i is not None and i < len(linha) else None

    respostas = []
    for linha in linhas[inicio:]:
        codigo = _texto(col(linha, "codigo"))
        if not codigo:
            continue
        respostas.append({
            "seq": _texto(col(linha, "seq")),
            "codigo": codigo,
            "valor": col(linha, "valor"),
            "marca": _texto(col(linha, "marca")),
        })
    return respostas


def aplicar_resposta(cot: Cotacao, fornecedor: str, respostas: list[dict]) -> dict:
    """Coloca os preços da resposta nas colunas do fornecedor.
    Casa pelo código do produto; se não achar, usa o SEQ.
    Retorna um resumo: quantos itens tiveram preço, quantos sem preço e
    códigos que não existem na cotação.
    """
    fornecedor = cot.adicionar_fornecedor(fornecedor)
    por_codigo = {_normalizar_codigo(it.codigo): it for it in cot.itens}
    por_seq = {it.seq: it for it in cot.itens}
    resumo = {"com_preco": 0, "sem_preco": 0, "nao_encontrados": []}
    for r in respostas:
        codigo = _normalizar_codigo(r["codigo"])
        it = por_seq.get(r["seq"])
        # O SEQ só vale se o código bate (o fornecedor pode ter reordenado).
        if it is None or _normalizar_codigo(it.codigo) != codigo:
            it = por_codigo.get(codigo) or (it if not codigo else None)
        if it is None:
            resumo["nao_encontrados"].append(r["codigo"])
            continue
        marca = r["marca"]
        if marca.upper() == "INFORMAR MARCA":
            marca = ""
        cot.definir_preco(fornecedor, it.seq, r["valor"], marca)
        if converter_valor(r["valor"]) is None:
            resumo["sem_preco"] += 1
        else:
            resumo["com_preco"] += 1
    cot.respondidos.add(fornecedor)
    return resumo


def adivinhar_fornecedor(caminho: str, fornecedores: list[str]) -> str:
    """Tenta achar o nome do fornecedor no nome do arquivo."""
    nome = os.path.basename(caminho).upper()
    for f in sorted(fornecedores, key=len, reverse=True):
        if f.replace(" ", "") in nome.replace(" ", "").replace("_", "").replace("-", ""):
            return f
    return ""


# ---------------------------------------------------------------- exportar

VERDE = PatternFill("solid", fgColor="C6EFCE")
AMARELO = PatternFill("solid", fgColor="FFEB9C")
CINZA = PatternFill("solid", fgColor="D9D9D9")


def exportar_excel(cot: Cotacao, caminho: str) -> None:
    """Gera um .xlsx com a aba COMPRAS preenchida e uma aba PEDIDOS com os
    itens separados pelo fornecedor de menor preço."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "COMPRAS"

    cab = ["SEQ", "NR_FABRICA", "ATENÇÃO MARCA", "QTDE", "PRODUTO"]
    for f in cot.fornecedores:
        cab += [f, f + " MARCA"]
    cab += ["MENOR", "EMPRESA", "2º MENOR", "EMPRESA", "PORCENTAGEM"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = CINZA
        c.alignment = Alignment(horizontal="center")

    base = 6
    for it in cot.itens:
        comp = cot.comparativo(it)
        linha = [it.seq, it.nr_fabrica, it.marca_exigida, it.qtd, it.descricao]
        for f in cot.fornecedores:
            p = cot.preco(f, it.seq)
            linha += [p.valor if p else None, p.marca if p else None]
        linha += [comp["menor"], comp["empresa"], comp["segundo"], comp["empresa2"], comp["diferenca"]]
        ws.append(linha)
        r = ws.max_row
        for i, f in enumerate(cot.fornecedores):
            cel = ws.cell(r, base + i * 2)
            cel.number_format = '"R$" #,##0.00'
            if comp["menor"] is not None and f == comp["empresa"]:
                cel.fill = VERDE
        n = base + len(cot.fornecedores) * 2
        ws.cell(r, n).number_format = '"R$" #,##0.00'
        ws.cell(r, n + 2).number_format = '"R$" #,##0.00'
        pct = ws.cell(r, n + 4)
        pct.number_format = "0.00%"
        if comp["diferenca"] is not None and comp["diferenca"] < LIMITE_DIFERENCA:
            pct.fill = AMARELO

    larguras = [6, 34, 22, 7, 44] + [11, 12] * len(cot.fornecedores) + [11, 13, 11, 13, 13]
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions

    wp = wb.create_sheet("PEDIDOS")
    wp.append(["FORNECEDOR", "SEQ", "CÓDIGO", "PRODUTO", "QTDE", "VALOR UNIT.", "MARCA", "TOTAL"])
    for c in wp[1]:
        c.font = Font(bold=True)
        c.fill = CINZA
    for f in cot.fornecedores:
        linhas = [(it, cot.preco(f, it.seq)) for it in cot.itens
                  if cot.comparativo(it)["empresa"] == f]
        if not linhas:
            continue
        total = 0.0
        for it, p in linhas:
            t = p.valor * (it.qtd or 0)
            total += t
            wp.append([f, it.seq, it.codigo, it.descricao, it.qtd, p.valor, p.marca, t])
        wp.append(["", "", "", f"TOTAL {f}", "", "", "", total])
        wp.cell(wp.max_row, 4).font = Font(bold=True)
        wp.cell(wp.max_row, 8).font = Font(bold=True)
        wp.append([])
    for linha in wp.iter_rows(min_row=2):
        linha[5].number_format = '"R$" #,##0.00'
        linha[7].number_format = '"R$" #,##0.00'
    for i, w in enumerate([14, 6, 22, 44, 7, 12, 14, 13], 1):
        wp.column_dimensions[get_column_letter(i)].width = w

    wb.save(caminho)
